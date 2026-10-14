import io
import pymupdf as fitz
from docx import Document
from pptx import Presentation
# 不受支持的后缀名
class UnsupportedFormatError(Exception):
    pass




def extract_segments(filename: str, data: bytes) -> list[tuple[str, dict | None]]:
    """按扩展名把上传文件转成 (文本, 定位) 段列表。

    定位是"能给就给"：PDF 按页拆，每页带 {"page_number": n}；
    其余格式以对应格式返回。
    返回多段而非一整段，是为了让切分在段内进行，PDF 的 chunk 不跨页，
    页码天然正确。
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "txt":
        return [(data.decode("utf-8", errors="ignore"), {
            "format": "txt",
        })]
    if ext == "md":
        text = data.decode("utf-8-sig", errors="replace")
        text = text.replace("\r\n", "\n").replace("\r", "\n")

        if not text.strip():
            return []

        return [
            (
                text,
                {
                    "format": "md",
                    "block_type": "paragraph",
                },
            )
        ]

    if ext == "pdf":
        with fitz.open(stream=data, filetype="pdf") as doc:
            return [
                (page.get_text(), {"format":"pdf","page_number": i + 1})
                for i, page in enumerate(doc)
            ]

    if ext == "docx":
        doc = Document(io.BytesIO(data))
        return [("\n".join(p.text for p in doc.paragraphs),  {
            "format": "docx"
        })]

    if ext == "pptx":
        prs = Presentation(io.BytesIO(data))
        segments = []
        for slide_number, slide in enumerate(prs.slides, start=1):
            # 标题占位符可能不存在
            title_shape = slide.shapes.title
            slide_title = title_shape.text.strip() if title_shape else ""
            parts = []
            for shape in slide.shapes:
                # 普通文本框
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        parts.append(text)
                # 表格：一行整理成一行文本
                elif shape.has_table:
                    for row in shape.table.rows:
                        cells = [
                            cell.text.strip()
                            for cell in row.cells
                        ]
                        if any(cells):
                            parts.append(" | ".join(cells))
            content = "\n".join(parts).strip()
            # 空白幻灯片不入库
            if not content:
                continue
            segments.append(
                (
                    content,
                    {
                        "format": "pptx",
                        "slide_number": slide_number,
                        "slide_title": slide_title or f"幻灯片 {slide_number}",
                    },
                )
            )
        return segments

    raise UnsupportedFormatError(ext)
